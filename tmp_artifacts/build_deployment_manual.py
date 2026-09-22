from __future__ import annotations

from datetime import date
from pathlib import Path
from urllib.parse import urlparse

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX_OUT = ROOT / "output" / "docx" / "AI_SPI_Inference_API_虛擬環境與離線AIPC部署手冊.docx"

FONT = "Heiti TC"
MONO = "Consolas"
NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F7FB"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
DARK = "1F2937"
GOLD = "8A6500"
PALE_GOLD = "FFF7DF"
RED = "9B1C1C"
PALE_RED = "FFF0F0"
GREEN = "176B45"
PALE_GREEN = "EAF7F0"
WHITE = "FFFFFF"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, *, name=FONT, size=None, color=DARK, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "cs"):
        rfonts.set(qn(f"w:{attr}"), name)
    rfonts.set(qn("w:eastAsia"), FONT if name == MONO else name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag in ("tblW", "tblInd", "tblLayout"):
        old = tbl_pr.find(qn(f"w:{tag}"))
        if old is not None:
            tbl_pr.remove(old)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[min(idx, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_para_shading(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_para_border_left(paragraph, color: str, size=14, space=8):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def add_field(paragraph, instruction: str, placeholder="1"):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MID_GRAY)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), FONT)
    r_pr.extend([r_fonts, color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_numbering_definition(doc: Document, *, bullet: bool) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, ind, spacing])
    lvl.extend([start, fmt, lvl_text, suff, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def set_num(paragraph, num_id: int):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    p_pr.append(num_pr)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.65)
section.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal.font.size = Pt(11)
normal.font.color.rgb = rgb(DARK)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Title", 30, NAVY, 0, 8),
    ("Subtitle", 14, MID_GRAY, 0, 8),
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, DARK_BLUE, 10, 5),
):
    style = styles[name]
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    style.font.bold = name.startswith("Heading") or name == "Title"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.0 if name in ("Title", "Subtitle") else 1.1
    style.paragraph_format.keep_with_next = True

bullet_num_id = add_numbering_definition(doc, bullet=True)
number_num_id = add_numbering_definition(doc, bullet=False)


def add_para(text="", *, bold_lead=None, italic=False, size=11, color=DARK, align=None, after=6, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_together = keep
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), size=size, color=color, bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]), size=size, color=color, italic=italic)
    else:
        set_run_font(p.add_run(text), size=size, color=color, italic=italic)
    return p


def add_bullet(text):
    p = doc.add_paragraph()
    set_num(p, bullet_num_id)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_run_font(p.add_run(text), size=10.5)
    return p


def add_step(text):
    p = doc.add_paragraph()
    set_num(p, number_num_id)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.25
    set_run_font(p.add_run(text), size=10.5)
    return p


def add_code(code: str):
    for line in code.strip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.08)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        set_para_shading(p, "F7F8FA")
        set_run_font(p.add_run(line if line else " "), name=MONO, size=8.2, color="263238")
    p.paragraph_format.space_after = Pt(7)


def add_callout(label: str, text: str, *, kind="info"):
    palette = {
        "info": (PALE_BLUE, BLUE),
        "warn": (PALE_GOLD, GOLD),
        "risk": (PALE_RED, RED),
        "ok": (PALE_GREEN, GREEN),
    }
    fill, accent = palette[kind]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.18
    set_para_shading(p, fill)
    set_para_border_left(p, accent)
    set_run_font(p.add_run(label + "  "), size=10.5, color=accent, bold=True)
    set_run_font(p.add_run(text), size=10.2, color=DARK)
    return p


def add_table(headers: list[str], rows: list[list[str]], widths: list[int], *, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), size=font_size, color=NAVY, bold=True)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            set_run_font(p.add_run(value), size=font_size, color=DARK)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    return table


def add_page_break():
    global _page_break_count
    _page_break_count += 1
    # Keep the cover isolated. Later sections flow naturally so a trailing line
    # cannot be stranded before an unconditional chapter break.
    if _page_break_count == 1:
        doc.add_page_break()


def add_heading(text: str, level=1):
    p = doc.add_heading(text, level=level)
    return p


def configure_footer_header():
    hdr = section.header
    p = hdr.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    set_run_font(p.add_run("AI SPI INFERENCE API  |  離線 AIPC 部署手冊"), size=8.5, color=MID_GRAY, bold=True)
    first_hdr = section.first_page_header
    first_hdr.paragraphs[0].text = ""
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""


configure_footer_header()
_page_break_count = 0

# Cover: editorial_cover pattern
for _ in range(4):
    add_para("", after=16)
p = add_para("部署與維運手冊", size=11, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
p.runs[0].bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
set_run_font(p.add_run("AI SPI Inference API"), size=30, color=NAVY, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
set_run_font(p.add_run("虛擬環境準備、離線 AIPC 搬運、設定、部署與疑難排解"), size=15, color=DARK_BLUE, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(54)
set_run_font(p.add_run("涵蓋 Conda、uv、Python venv 與 Windows 服務化部署"), size=11, color=MID_GRAY, italic=True)

add_table(
    ["文件屬性", "內容"],
    [
        ["專案", "ai-spi-inference-api 0.1.0"],
        ["主要目標", "無網際網路的 NVIDIA AIPC"],
        ["Python", ">= 3.12, < 3.13"],
        ["文件日期", "2026-08-31"],
        ["建議路徑", "uv 鎖定檔 + 可攜 uv 快取；備援為 Conda Pack"],
    ],
    [1875, 7485],
    font_size=9.5,
)
add_para("本手冊以專案目前的 README、pyproject.toml、uv.lock、Windows 啟動器、設定驗證程式與環境檢查腳本為準。", size=9.2, color=MID_GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)

add_page_break()

add_heading("如何使用本手冊", 1)
add_callout("先看這裡", "如果只需要一次成功部署，採用第 4 章的「建議方案 A：uv 離線快取」。若 AIPC 不允許安裝 Python 或 uv，改用第 7 章的 Conda Pack。不要直接複製現有 .venv。", kind="ok")
add_heading("閱讀導覽", 2)
for text in [
    "第 1-3 章：確認專案架構、版本與部署前提。",
    "第 4-7 章：分別準備 uv、venv、Conda 與離線移轉包。",
    "第 8-11 章：在 AIPC 設定路徑、模型、TensorRT、啟動器與 Windows 服務。",
    "第 12-13 章：驗收、上線、備份與回復。",
    "第 14 章：依症狀查找 troubleshooting。",
]:
    add_bullet(text)
add_heading("範圍與假設", 2)
add_bullet("遠端 AIPC 沒有網際網路，但可透過核准的 USB、內網檔案站或其他離線媒體取得部署包。")
add_bullet("AIPC 為 Windows 10/11 或 Windows Server x86-64，配備 NVIDIA GPU；Linux 可沿用依賴原則，但批次檔與 NSSM 章節不適用。")
add_bullet("SPI 機台共享資料夾可能仍透過廠內 LAN/SMB 提供；「離線」在此指無外網。若連 LAN 也沒有，watch_root 與輸出路徑需改成本機交換目錄。")
add_bullet("模型權重不在目前 Git 檔案清單中，部署包必須另帶 .pt、.onnx 或 .engine 檔與其校驗碼。")
add_callout("重要限制", "Python venv 官方定義為不可搬移、應在目標位置重建；TensorRT engine 也通常綁定平台、TensorRT 版本與 GPU。這兩類資產都不能假設「複製過去就一定能跑」。", kind="warn")

add_heading("1. 專案部署輪廓", 1)
add_para("正式部署建議使用三階段 durable pipeline。Ingest、Inference、Publisher 是獨立常駐程序，以 AIPC 本機 SQLite WAL 協調；模型服務則在 8000 與 8002 提供推論。相容 API 5050 是選配，不是三階段流程的必要元件。")
add_table(
    ["元件", "連接埠", "責任", "生產必要性"],
    [
        ["Publisher", "-", "到截止點接管、先發布 Primary、必要時全 23", "必要，且先啟動"],
        ["PatchCore anomaly", "8000", "輸出 anomaly_score", "預設 required"],
        ["Distance detection", "8002", "輸出 min_pad_distance", "預設 required"],
        ["Inference", "-", "以最早截止期限執行模型並寫 manifest", "必要"],
        ["Ingest", "-", "監看、驗證、備份原始資料並排入 READY", "必要，最後啟動"],
        ["Paste detection", "8001", "paste_pixels；較重", "預設停用"],
        ["Compatibility API", "5050", "/process、/health、/ready", "選配"],
    ],
    [1800, 900, 4680, 1980],
    font_size=8.8,
)
add_para("啟動順序：Publisher -> 模型 8000/8002 且 health=healthy -> Inference -> Ingest。這不是批次檔編號順序；Publisher 必須先在線，因為它負責 deadline fallback。", bold_lead="啟動順序：")

add_page_break()

add_heading("2. 版本基準與不可省略的系統元件", 1)
add_table(
    ["層級", "專案基準", "離線包必備"],
    [
        ["Python", ">=3.12,<3.13；啟動器只接受 3.12", "Python 3.12 x64 安裝程式或已封裝 Conda 環境"],
        ["PyTorch", "2.7.1 + CUDA 12.8 index", "torch 2.7.1+cu128 與 torchvision 0.22.1+cu128"],
        ["ONNX Runtime", "onnxruntime-gpu 1.19.0", "相符 wheel 與 CUDA/cuDNN 執行庫"],
        ["PyCUDA", "2024.1", "預建 wheel，或離線編譯工具鏈"],
        ["TensorRT", "10.x，但專案拒絕 10.1.x", "與 driver/CUDA 相符的 NVIDIA 安裝包及 Python wheel"],
        ["uv", "pyproject.toml + uv.lock", "相同版本 uv.exe、uv.lock、已暖機快取"],
        ["Windows runtime", "本機服務與原生 wheel", "NVIDIA driver、VC++ Redistributable、必要時 VS Build Tools"],
    ],
    [1800, 2880, 4680],
    font_size=8.9,
)
add_callout("目前工作區觀察", "專案根目錄存在 uv.lock，且 requires-python=3.12；但 uv.lock 目前被 .gitignore 忽略，現有 .venv 則是 Python 3.10。部署包請明確包含 uv.lock，並在 AIPC 重新建立 Python 3.12 的 .venv。", kind="risk")
add_heading("2.1 AIPC 上線前資料表", 2)
for item in [
    "作業系統版本與更新層級、CPU 架構（x86-64）、GPU 型號與 compute capability。",
    "NVIDIA driver 版本、CUDA runtime/toolkit、cuDNN、TensorRT 完整版本與 Python binding 版本。",
    "AIPC 專案固定路徑，例如 D:\\Dre\\JQ_SPI_02_AI_API；服務帳號與其檔案權限。",
    "SPI watch_root、machine return 目錄、備份目錄、本機 state/result/log 目錄。",
    "模型檔名、格式、SHA256、input shape、batch profile、前處理契約與來源。",
    "允許的輸入輸出媒體、惡意程式掃描流程、離線包簽核人與回復版本。",
]:
    add_bullet(item)

add_heading("3. 三種虛擬環境方案如何選", 1)
add_table(
    ["方案", "優點", "限制", "本專案建議"],
    [
        ["uv", "快、使用 uv.lock、與 pyproject.toml 一致", "離線端必須有完整快取與同版 uv", "首選；在線同平台暖機後搬移 cache"],
        ["venv + wheelhouse", "只依賴標準 Python/pip、可稽核 wheel", "需自行解決 PyCUDA/TensorRT 等原生套件", "適合嚴格軟體白名單"],
        ["Conda + conda-pack", "可把 Python 與多數套件一起封裝", "包大、同 OS/架構、首次解包後不可再搬", "AIPC 無 Python/uv 時的可靠備援"],
    ],
    [1200, 2520, 2520, 3120],
    font_size=8.7,
)
add_callout("建議決策", "正式包以 uv 為主、Conda Pack 為備援；wheelhouse 同時保留做套件級修復。三者都要在與 AIPC 相同 OS、架構、Python minor 與 GPU 軟體基準的「在線孿生機」產生。", kind="ok")

add_page_break()

add_heading("4. 建議方案 A：用 uv 準備可離線部署包", 1)
add_heading("4.1 在線 Windows 孿生機準備", 2)
add_step("安裝與 AIPC 相同的 Python 3.12 x64、NVIDIA driver/CUDA/TensorRT 基準；使用與目標端完全相同的專案路徑更容易重現問題。")
add_step("取得固定版本 uv.exe，建立獨立部署快取目錄，避免混入開發者其他專案的套件。")
add_code(r"""
cd /d D:\build\ai_spi_inference_api
set UV_CACHE_DIR=D:\offline_bundle\uv-cache
uv --version
uv lock --check
uv venv --python 3.12 --clear
uv sync --locked --extra cuda --no-group dev
""")
add_step("如需在 AIPC 轉模型，額外同步 tensorrt-export；如需測試工具，再加入 dev。預設停用 paste，不要為生產核心包引入 Git 依賴。")
add_code(r"""
uv sync --locked --extra cuda --extra tensorrt-export --no-group dev
rem 僅驗證包才加：uv sync --locked --extra cuda --group dev
""")
add_step("執行環境、GPU、provider 與測試驗證；成功後才封裝。")
add_code(r"""
uv run python scripts\check_env.py
uv run python -c "import torch; print(torch.__version__, torch.version.cuda, torch.cuda.is_available())"
uv run python -c "import onnxruntime as ort; print(ort.get_available_providers())"
uv run pytest tests\unit tests\integration
""")
add_callout("快取完整性", "uv --offline 只會使用本機檔案與快取。要確認真的完整，請在孿生機先中斷外網或封鎖網路，另建空白 .venv，使用相同 UV_CACHE_DIR 執行第 4.3 節命令。只做一次在線 uv sync 不等於已驗證可離線重建。", kind="warn")

add_heading("4.2 部署包目錄建議", 2)
add_code(r"""
AI_SPI_OFFLINE_2026-08-31\
  project\                 專案程式、uv.lock、設定範本、批次檔
  uv\uv.exe                與暖機時完全相同版本
  uv-cache\                專用且經離線重建驗證的 uv cache
  installers\              Python 3.12、driver、VC++、CUDA、TensorRT、NSSM
  wheelhouse\              修復用 wheels；含 CUDA torch 與原生套件
  models\                  模型、engine.json、SHA256SUMS.txt
  config\ai_server.aipc.json
  MANIFEST.txt
  SHA256SUMS.txt
""")
add_bullet("不要打包 .venv、log、SQLite state、pipeline_results、既有輸出或其他機台的 config/ai_server.json。")
add_bullet("一定要明確帶入 uv.lock；本工作區的 .gitignore 會忽略它，單純 git clone/zip tracked files 可能漏掉。")
add_bullet("所有檔案計算 SHA256；搬到 AIPC 後先驗證，再安裝或解壓。")

add_heading("4.3 AIPC 離線重建 uv 環境", 2)
add_step("將部署包複製到 AIPC 本機磁碟，驗證 SHA256；先安裝 driver、VC++、Python 3.12、CUDA/TensorRT。安裝後重開機。")
add_step("把 project 放到固定路徑，uv.exe 與 uv-cache 放本機磁碟；不要從 USB 或 SMB 直接執行。")
add_code(r"""
cd /d D:\Dre\JQ_SPI_02_AI_API
set UV_CACHE_DIR=D:\Dre\AI_SPI_OFFLINE\uv-cache
set UV_OFFLINE=1
D:\Dre\AI_SPI_OFFLINE\uv\uv.exe venv --python C:\Python312\python.exe --clear
D:\Dre\AI_SPI_OFFLINE\uv\uv.exe sync --frozen --offline --extra cuda --no-group dev --no-python-downloads
""")
add_para("若 AIPC 也負責 TensorRT 轉換，將上一行加入 --extra tensorrt-export。若服務只跑既有 .engine，仍需 TensorRT runtime/Python binding，但不需要 onnx/onnxslim。")
add_step("固定啟動器使用的 Python，避免 PATH 撿到錯誤版本。")
add_code(r"""
set PYTHON_EXE=D:\Dre\JQ_SPI_02_AI_API\.venv\Scripts\python.exe
%PYTHON_EXE% --version
%PYTHON_EXE% scripts\check_env.py
""")

add_page_break()

add_heading("5. 方案 B：標準 venv + wheelhouse", 1)
add_para("這是最透明的離線安裝方式：AIPC 先離線安裝官方 Python 3.12，再就地建立 .venv，所有套件只從經核准的 wheelhouse 安裝。官方 Python 文件明確指出 venv 不應被視為可搬移資產，因此不要從孿生機複製 .venv。")
add_heading("5.1 在線孿生機建立 wheelhouse", 2)
add_step("在完全相同的 Windows/Python/GPU 基準完成一次在線安裝；輸出實際解析版本。")
add_code(r"""
py -3.12 -m venv D:\build\spi-wheel-build
D:\build\spi-wheel-build\Scripts\activate.bat
python -m pip install --upgrade pip wheel
python -m pip install torch==2.7.1 torchvision==0.22.1 --index-url https://download.pytorch.org/whl/cu128
python -m pip install -r requirements-cuda.txt
python -m pip freeze > offline_bundle\requirements-resolved-win-py312.txt
""")
add_step("把解析後的所有套件下載/建成 wheel。PyCUDA 若沒有相容 wheel，必須在具備 CUDA headers 與 MSVC Build Tools 的孿生機先建好，不要把 sdist 留到離線 AIPC 才編譯。")
add_code(r"""
python -m pip download -d offline_bundle\wheelhouse ^
  -r offline_bundle\requirements-resolved-win-py312.txt ^
  --extra-index-url https://download.pytorch.org/whl/cu128
python -m pip wheel -w offline_bundle\wheelhouse pycuda==2024.1
""")
add_step("另外放入 TensorRT 安裝包/whl；它未固定在 requirements-cuda.txt，需依 AIPC 的 NVIDIA 軟體矩陣選版。")
add_callout("供應鏈注意", "混用 PyPI 與額外 index 時，應在受控孿生機完成解析並保留檔案雜湊。離線 AIPC 一律使用 --no-index，避免任何未核准來源。", kind="warn")

add_heading("5.2 AIPC 離線安裝", 2)
add_code(r"""
cd /d D:\Dre\JQ_SPI_02_AI_API
py -3.12 -m venv .venv
.venv\Scripts\python.exe -m pip install --no-index ^
  --find-links D:\Dre\AI_SPI_OFFLINE\wheelhouse ^
  -r D:\Dre\AI_SPI_OFFLINE\requirements-resolved-win-py312.txt
.venv\Scripts\python.exe -m pip check
.venv\Scripts\python.exe scripts\check_env.py
""")
add_para("若 PowerShell 不允許 Activate.ps1，不必調降全機安全政策；直接使用 .venv\\Scripts\\python.exe 呼叫即可。所有專案批次檔也支援用 PYTHON_EXE 指定完整路徑。")

add_heading("6. 方案 C：Conda 在線建立環境", 1)
add_para("專案附有 setup.bat，會建立 py312_cu128_j15 並在線下載 PyTorch CUDA 與 requirements.txt。它適合有網路的建置機，不適合直接在離線 AIPC 執行。")
add_code(r"""
cd /d D:\build\ai_spi_inference_api
setup.bat
conda run -n py312_cu128_j15 python scripts\check_env.py
conda run -n py312_cu128_j15 python -m pip check
""")
add_callout("版本一致性", "requirements.txt 是保留給 Conda/plain-pip 的 legacy mirror；依賴主來源仍是 pyproject.toml 與 uv.lock。建置前先比較差異，避免 Conda 包與 uv 生產包產生不同版本。", kind="info")

add_heading("7. Conda Pack：把環境帶到無 Python 的 AIPC", 1)
add_heading("7.1 在線孿生機封裝", 2)
add_code(r"""
conda install -n base -c conda-forge conda-pack -y
conda pack -n py312_cu128_j15 ^
  -o D:\offline_bundle\py312_cu128_j15.zip ^
  --format zip
""")
add_bullet("建置與目標必須是相同 OS 與架構；Windows 包不能搬到 Linux。")
add_bullet("封裝前先在該環境驗證 torch/CUDA、ONNX provider、TensorRT import 與專案測試。")
add_bullet("Conda Pack 能處理環境前綴搬移，但 NVIDIA driver 仍屬系統層，不能只靠環境包。")

add_heading("7.2 AIPC 解包與固定路徑", 2)
add_code(r"""
powershell -NoProfile -Command ^
  "Expand-Archive -Path D:\AI_SPI_OFFLINE\py312_cu128_j15.zip -DestinationPath D:\Dre\envs\py312_cu128_j15"
D:\Dre\envs\py312_cu128_j15\Scripts\activate.bat
D:\Dre\envs\py312_cu128_j15\Scripts\conda-unpack.exe
set PYTHON_EXE=D:\Dre\envs\py312_cu128_j15\python.exe
%PYTHON_EXE% scripts\check_env.py
""")
add_callout("解包後不要移動", "conda-unpack 完成後，該環境視為固定在目前路徑。若位置錯誤，刪除這個解包副本並從原始 zip 重新解到正確位置，不要再搬資料夾。", kind="warn")

add_page_break()

add_heading("8. 專案、模型與設定的離線移轉", 1)
add_heading("8.1 專案包應包含與排除", 2)
add_table(
    ["類別", "包含", "排除/另行處理"],
    [
        ["程式", "app、scripts、tests、*.py、*.bat、README、pyproject.toml、uv.lock", ".git、cache、開發機 .venv"],
        ["設定", "ai_server.example.json、AIPC 專用 config", "其他機台的實際路徑與密碼"],
        ["模型", "PatchCore、center、pad、可選 paste；engine.json", "來源不明或未驗證 artifact"],
        ["運行資料", "空白 state/log/result 目錄或由程式建立", "舊 SQLite、舊 WAL、舊 job/result"],
        ["系統安裝包", "driver、VC++、Python/Conda、CUDA、TensorRT、uv、NSSM", "與 AIPC 硬體不相符版本"],
    ],
    [1500, 4320, 3540],
    font_size=8.8,
)
add_heading("8.2 SHA256 與部署清單", 2)
add_code(r"""
rem 在建置機產生（PowerShell）
Get-ChildItem D:\offline_bundle -Recurse -File |
  Get-FileHash -Algorithm SHA256 |
  ForEach-Object { "$($_.Hash) *$($_.Path)" } |
  Set-Content D:\offline_bundle\SHA256SUMS.txt
""")
add_para("在 AIPC 逐項重新計算並比對。MANIFEST 至少記錄：包版號、Git commit、uv 版本、Python、torch/CUDA、onnxruntime providers、TensorRT、driver、GPU、模型 SHA256、設定檔 SHA256、建置與驗收人員。")
add_callout("信任邊界", "TensorRT engine 與 Python pickle 型 .pt/.pth/.ckpt 都可能執行或載入可執行內容。只使用自建或經核准來源；PatchCore .pt 只有在可信時才加 --trust-pickle。", kind="risk")

add_heading("9. AIPC 專案設定", 1)
add_heading("9.1 建立 AIPC 專用設定檔", 2)
add_step("以 config\\ai_server.example.json 複製出機台專用檔，例如 D:\\Dre\\config\\ai_server.aipc01.json；不要直接覆蓋共享範本。")
add_step("以 AI_CONFIG_PATH 指向它，並讓所有模型/worker/相容 API 使用同一份經簽核設定。修改後必須重啟各程序，因為設定在每個 process 內快取。")
add_code(r"""
set AI_CONFIG_PATH=D:\Dre\config\ai_server.aipc01.json
set PYTHON_EXE=D:\Dre\JQ_SPI_02_AI_API\.venv\Scripts\python.exe
""")
add_table(
    ["設定欄位", "部署原則", "常見錯誤"],
    [
        ["paths.external_output_root", "機台可見的 Primary 輸出；確認服務帳號可寫", "NSSM 看不到使用者映射磁碟"],
        ["paths.backup_output_root", "AIPC 本機持久備份", "放在 SMB 導致 SLA/可靠性降低"],
        ["pipeline.watch_root", "SPI timestamp 來源；durable Ingest 只看這個欄位", "誤改頂層 legacy watch_root"],
        ["pipeline.database_path", "AIPC 本機 NTFS；所有 worker 共用同一檔", "放 UNC/SMB 造成 WAL 鎖定"],
        ["pipeline.staging_root", "建議 null，直接用不可變 raw backup", "多複製一次吃掉 30 秒預算"],
        ["pipeline.result_root", "AIPC 本機暫存 manifest/processed", "目錄無寫入權限"],
        ["model_clients", "8000/8002 enabled+required；URL/欄位唯一", "required service 未健康導致全 23"],
        ["output", "durable 必須 is_pass_only + machine_return", "其他模式導致啟動驗證失敗"],
        ["reliability", "deadline 30、reserve 5；lease 必須符合驗證關係", "reserve/lease 關係不合法"],
        ["logging.log_dir", "本機可寫；依 stage 產生 system.*", "用相對路徑卻誤判位置"],
    ],
    [2700, 3780, 2880],
    font_size=8.35,
)

add_heading("9.2 設定驗證命令", 2)
add_code(r"""
%PYTHON_EXE% -c "from pathlib import Path; from app.core.config import load_config; c=load_config(Path(r'%AI_CONFIG_PATH%')); print(c.model_dump_json(indent=2))"
%PYTHON_EXE% -m app.pipeline publisher --once --config "%AI_CONFIG_PATH%"
""")
add_para("第二行同時做 Publisher 所需路徑的 preflight；請先使用測試設定與空白測試目錄。若正式資料夾已有待處理工作，不要用 --once 做隨意測試。")

add_heading("9.3 建議的路徑與權限", 2)
add_code(r"""
D:\Dre\JQ_SPI_02_AI_API\        程式與 .venv
D:\Dre\config\                  機台設定
D:\Dre\JQ_SPI_02_AI_API\state\ SQLite + WAL
D:\Dre\JQ_SPI_02_AI_API\pipeline_results\
D:\Dre\JQ_SPI_02_AI_API\log\
D:\Dre\JQ_SPI_02_AI_API\backup\
""")
add_bullet("服務帳號對專案程式至少需讀取；對 state/result/log/backup 與 machine return 需修改/建立/重新命名權限。")
add_bullet("若 external_output_root/watch_root 是 SMB，服務優先使用 UNC 並指定具權限的服務帳號；Windows service 通常看不到互動登入者的映射磁碟機。")
add_bullet("不要把 SQLite database_path 放在 UNC/SMB；它必須在 AIPC 本機。")

add_page_break()

add_heading("10. 模型與 TensorRT", 1)
add_heading("10.1 選擇推論格式", 2)
add_table(
    ["格式", "後端", "AIPC 注意事項"],
    [
        [".pt", "PyTorch / Ultralytics", "需 CUDA torch；PatchCore 可信 pickle 才可 --trust-pickle"],
        [".onnx", "ONNX Runtime GPU", "available providers 必須包含 CUDAExecutionProvider；YOLO batch 動態或固定 8"],
        [".engine", "TensorRT/PyCUDA", "最快但最綁定；建議在目標 AIPC 建立並驗證"],
    ],
    [1200, 2760, 5400],
    font_size=8.9,
)
add_heading("10.2 模型服務環境變數", 2)
add_code(r"""
set SPI_MODEL_DEVICE=cuda:0
set PATCHCORE_MODEL_PATH=models\patchcore\model_fp16.engine
set DISTANCE_CENTER_MODEL_PATH=models\distance\center.engine
set DISTANCE_PAD_MODEL_PATH=models\distance\pad.engine
set PATCHCORE_MODEL_ARGS=
start_model_services.bat
""")
add_para("相對模型路徑以專案根目錄為基準；NSSM 的 AppDirectory 必須設為專案根目錄。若用可信 PatchCore .pt，再設定 PATCHCORE_MODEL_ARGS=--trust-pickle ... 並完整對齊訓練前處理。")
add_heading("10.3 在目標 AIPC 轉 TensorRT", 2)
add_para("TensorRT plan 預設只相容建立它的 TensorRT 版本、平台與 GPU 類型；本專案也明確要求在目標 NVIDIA AIPC 建 engine，且拒絕 TensorRT 10.1.x。輸入 artifact、profile、前處理與 batch 必須保留在 engine.json。")
add_code(r"""
%PYTHON_EXE% -m pip install --no-index --find-links D:\AI_SPI_OFFLINE\wheelhouse ^
  -r requirements-tensorrt-export.txt

convert_to_tensorrt.bat yolo ^
  --input models\distance\pad.pt ^
  --output models\distance\pad.engine ^
  --task detect --precision fp16 --batch 8 --imgsz 640 640 ^
  --workspace-gib 4 --save-onnx models\distance\pad.onnx
""")
add_bullet("center 與 pad 的 YOLO service 會以 batch 8 執行；動態 ONNX 必須提供涵蓋 1/最佳/8 的 profile，靜態圖則固定 8。")
add_bullet("PatchCore 要對齊 NCHW RGB float32 [0,1]、input size、center crop、ImageNet/custom normalization 與 score output。")
add_bullet("轉換後用代表性 SPI 影像比較 PyTorch/ONNX/TensorRT 分數與最終 22/23 決策，不只看 engine 能否載入。")

add_heading("11. 部署與服務化", 1)
add_heading("11.1 第一次手動啟動", 2)
add_step("設定 PYTHON_EXE、AI_CONFIG_PATH 與模型路徑。先啟動 03 Publisher。")
add_step("執行 start_model_services.bat，輪詢 8000 與 8002；HTTP 200 不夠，JSON 內 status 必須是 healthy，且 backend/model/device 正確。")
add_step("模型健康後啟動 02 Inference；最後啟動 01 Ingest。")
add_code(r"""
start "SPI Publisher" 03_pipeline_publisher.bat
start_model_services.bat
powershell -NoProfile -Command "Invoke-RestMethod http://127.0.0.1:8000/health | ConvertTo-Json -Depth 6"
powershell -NoProfile -Command "Invoke-RestMethod http://127.0.0.1:8002/health | ConvertTo-Json -Depth 6"
start "SPI Inference" 02_pipeline_inference.bat
start "SPI Ingest" 01_pipeline_ingest.bat
""")
add_callout("不要混跑", "Durable mode 使用三個 app.pipeline worker；Compatibility mode 才用 scan_jobs.py + app.main。不要讓 legacy scanner 與 Stage 01 同時監看同一 share，否則可能重複發布。", kind="risk")

add_heading("11.2 NSSM 服務化", 2)
add_para("每個 worker 與模型各自建立服務，才能獨立重啟。以下示例省略模型參數；實際請把 model path/device 放入 AppParameters 或 AppEnvironmentExtra。")
add_code(r"""
set REPO=D:\Dre\JQ_SPI_02_AI_API
set PYTHON_EXE=%REPO%\.venv\Scripts\python.exe
set CONFIG=D:\Dre\config\ai_server.aipc01.json

nssm install SPI_03_Publisher "%PYTHON_EXE%" -m app.pipeline publisher
nssm set SPI_03_Publisher AppDirectory "%REPO%"
nssm set SPI_03_Publisher AppEnvironmentExtra "AI_CONFIG_PATH=%CONFIG%"
nssm set SPI_03_Publisher Start SERVICE_AUTO_START

nssm install SPI_02_Inference "%PYTHON_EXE%" -m app.pipeline inference
nssm set SPI_02_Inference AppDirectory "%REPO%"
nssm set SPI_02_Inference AppEnvironmentExtra "AI_CONFIG_PATH=%CONFIG%"

nssm install SPI_01_Ingest "%PYTHON_EXE%" -m app.pipeline ingest
nssm set SPI_01_Ingest AppDirectory "%REPO%"
nssm set SPI_01_Ingest AppEnvironmentExtra "AI_CONFIG_PATH=%CONFIG%"
""")
add_bullet("服務帳號需能讀寫 watch/output/backup/state/result/log；若使用 UNC，請用明確帳號，不要依賴登入者映射磁碟。")
add_bullet("設定 stdout/stderr 旋轉與故障自動重啟；應用程式本身另寫 system.ingest、system.inference、system.publisher。")
add_bullet("開機自動啟動不代表依賴已就緒。先啟動 Publisher，再讓模型健康後啟動 Inference/Ingest；可用 delayed start 或外部健康檢查腳本。")

add_page_break()

add_heading("12. 上線前驗收", 1)
add_heading("12.1 環境與 GPU", 2)
for item in [
    "[ ] %PYTHON_EXE% --version 為 3.12.x，sys.executable 指向預期環境。",
    "[ ] scripts/check_env.py 所有 required base package 為 OK。",
    "[ ] torch.__version__ 為 2.7.1 CUDA build；torch.cuda.is_available() 為 True。",
    "[ ] onnxruntime providers 含 CUDAExecutionProvider（若使用 ONNX GPU）。",
    "[ ] tensorrt 與 pycuda 可 import（若使用 engine），driver/CUDA/TensorRT 版本已記錄。",
    "[ ] pip check 或 uv sync --locked --offline 可成功重建。",
]:
    add_bullet(item)
add_heading("12.2 設定、目錄與服務", 2)
for item in [
    "[ ] AI_CONFIG_PATH 指向 AIPC 專用檔；Pydantic load_config 驗證成功。",
    "[ ] SQLite、backup、result、log 在 AIPC 本機；watch/output 權限以服務帳號驗證。",
    "[ ] 8000/8002 health 回報 status=healthy、模型路徑與 device 正確。",
    "[ ] Stage 03、模型、Stage 02、Stage 01 依正確順序啟動，沒有 legacy scanner 競爭。",
    "[ ] 所有模型與部署檔 SHA256 與 MANIFEST 一致。",
]:
    add_bullet(item)
add_heading("12.3 端到端測試", 2)
add_step("先以隔離的測試 config、watch_root 與 output_root 執行，不要直接對正式 SPI share。")
add_step("放入一個符合今日 14 位 timestamp 的小型 job；確認 CSV 欄位與所有期待影像存在、非空、可由 OpenCV 解碼。")
add_step("確認 SQLite 狀態依序到 READY / INFERENCING / RESULT_READY / PRIMARY_RETURNED / DONE。")
add_step("確認 Primary 先出現在 external_output_root，只有 is_pass 改變；原始編碼、delimiter、欄位順序、quoting、line ending 保留。")
add_step("確認 backup 下有 immutable raw、ai_result/returned、processed 與 manifest.json；log 記錄 latency 與 deadline_met。")
add_step("故意停止一個 required model，重跑隔離測試，確認 reason 為 required_model_failure/timeout 或 publish_reserve_reached，所有列安全回傳 23。")
add_callout("驗收門檻", "正常路徑與 fail-safe 路徑都必須通過。這個系統的安全承諾不只是「AI 有結果」，也包含 required model 失敗或超時時仍能以全 23 發布。", kind="ok")

add_heading("13. 上線、備份與回復", 1)
add_heading("13.1 正式切換", 2)
add_step("凍結程式、uv.lock、模型與 config；記錄 SHA256 與版本。")
add_step("停止舊 scanner/舊 API 對相同 share 的自動處理；確認沒有殘留程序。")
add_step("備份舊版專案、環境、config、模型與必要 SQLite；不要覆寫唯一可回復版本。")
add_step("按 Publisher -> models -> Inference -> Ingest 啟動；監控首批 job 與 deadline。")
add_heading("13.2 回復策略", 2)
add_bullet("程式回復：切回上一個只讀版本目錄與其配套 config/模型/環境，不在原目錄混裝套件。")
add_bullet("模型回復：.engine 必須與上一版 TensorRT/GPU 基準成套；不可只換單一 engine。")
add_bullet("狀態回復：先停所有三個 worker，再處理 SQLite/WAL；不要在程序執行時手工複製或刪改資料庫。")
add_bullet("Primary 已發布的 job 具 at-least-once 行為；重啟可能以相同決策覆寫已回傳 CSV，這是設計內行為。")
add_bullet("保留 manifest 與 stage logs，因 SQLite 是 job state，不是完整 audit event ledger。")

add_page_break()

add_heading("14. Troubleshooting", 1)
add_para("先依序取得四類證據：目前使用的 Python、環境/GPU/provider、設定驗證、stage/model logs。不要先重裝全部套件，否則會抹去版本與錯誤線索。")
add_code(r"""
where python
echo %PYTHON_EXE%
%PYTHON_EXE% --version
%PYTHON_EXE% scripts\check_env.py
nvidia-smi
""")

troubles = [
    (
        "14.1 啟動器顯示 Python 3.12 is required",
        "症狀：resolve_python.bat 找到 .venv，但立即退出；或 uv 說 requires-python 不相容。",
        "原因：環境不是 3.12。此工作區目前的 .venv 實際為 Python 3.10，是可重現的典型情況。",
        "處理：停止使用該環境，改名保留後以 Python 3.12 重建；或明確 set PYTHON_EXE 指向正確環境。不要把舊 .venv 複製到 AIPC。",
    ),
    (
        "14.2 uv --offline 找不到套件或仍嘗試連線",
        "症狀：cache miss、No solution found、failed to fetch，或離線重建失敗。",
        "原因：uv-cache 未含目標 platform marker 的 wheel、uv.lock 漏包、uv 版本不同、Python/OS/架構不同，或 paste 的 Git 依賴未預先快取。",
        "處理：回到相同 Windows x86-64 孿生機，指定乾淨 UV_CACHE_DIR，使用相同 uv.exe 與 uv.lock 完成一次真正斷網重建；生產包不要啟用 paste，除非已把 Git commit 固定並完整快取/建 wheel。",
    ),
    (
        "14.3 pip --no-index 顯示 No matching distribution",
        "症狀：某套件只有 sdist，或 wheel tag 不符 cp312/win_amd64。",
        "原因：wheelhouse 在 macOS/Linux 或不同 Python 建立；PyCUDA/TensorRT 等原生套件沒有適配 wheel。",
        "處理：在 Windows Python 3.12 x64 孿生機重新下載/建 wheel；用 pip debug --verbose 檢查 tag。PyCUDA 需 CUDA headers/MSVC；TensorRT wheel 從相符 NVIDIA 安裝包取得。",
    ),
    (
        "14.4 DLL load failed / 找不到指定模組",
        "症狀：import torch、onnxruntime、tensorrt、pycuda 或 cv2 失敗。",
        "原因：VC++ runtime、CUDA/cuDNN/TensorRT DLL 缺失，PATH 順序錯，或 32/64 位元不一致。",
        "處理：檢查 nvidia-smi、where 對應 DLL、系統 PATH 與安裝版本；重開機後再測。不要靠把不明 DLL 複製進專案目錄解決。",
    ),
    (
        "14.5 torch.cuda.is_available() = False",
        "症狀：torch 可 import，但模型使用 cuda:0 失敗。",
        "原因：裝到 CPU wheel、driver 不支援、GPU 不可見，或服務帳號環境不同。",
        "處理：確認 torch.__version__ 含 +cu128、torch.version.cuda、nvidia-smi；以實際服務帳號執行相同命令。必要時從核准 cu128 wheel 重建環境。",
    ),
    (
        "14.6 ONNX Runtime 沒有 CUDAExecutionProvider",
        "症狀：available providers 只有 CPUExecutionProvider；專案在 cuda:N 啟動時拒絕默默回 CPU。",
        "原因：裝到 onnxruntime 而非 onnxruntime-gpu，或 CUDA/cuDNN DLL 不相容。",
        "處理：移除 CPU/GPU 套件衝突，離線安裝 onnxruntime-gpu==1.19.0，重查 provider 與 native DLL。不要以 CPU fallback 當作正式修復。",
    ),
    (
        "14.7 TensorRT engine 無法反序列化",
        "症狀：版本、magic tag、compute capability、binding/profile 或 engine validity 錯誤。",
        "原因：engine 在其他 OS/GPU/TensorRT 建立、runtime 升級、檔案損壞，或用到 10.1.x。",
        "處理：核對 engine.json 與 SHA256；在目標 AIPC 用核准的 TensorRT 10.x（非 10.1.x）重新轉換，確認 batch/profile 與 preprocessing。",
    ),
    (
        "14.8 模型 /health 回 200，但還不能處理",
        "症狀：HTTP 成功，JSON status 卻是 initializing/unhealthy，或模型資訊不符。",
        "原因：TensorRT engine 仍初始化、模型路徑錯、GPU OOM，或 backend 載入失敗。",
        "處理：以 payload 的 status=healthy 為啟動門檻；查看模型視窗/stderr，確認 model path、device、free VRAM，再啟動 Stage 02/01。",
    ),
    (
        "14.9 Worker 一啟動就退出 / Config invalid",
        "症狀：pipeline.enabled、路徑 preflight 或 Pydantic validation error。",
        "原因：watch_root 空白、沒有 required model、duplicate name/target_column、deadline/reserve/lease 關係不合法，或 output 模式不符 durable contract。",
        "處理：用第 9.2 節 load_config 命令讀完整錯誤；durable 必須 is_pass_only + machine_return，worker_lease > deadline，publisher heartbeat < lease < reserve。",
    ),
    (
        "14.10 Ingest 一直沒有 READY",
        "症狀：timestamp folder 被反覆掃描但不進 READY。",
        "原因：不是今日 14 位資料夾、settle 期間持續變動、CSV 無資料或缺欄、期待影像缺失/空檔/不可解碼、影像 key 重複。",
        "處理：查看 system.ingest；核對 component_name、Array_id、Pad_id 與 image_name_template。歷史資料用 --once --date YYYY-MM-DD，並在測試設定執行。",
    ),
    (
        "14.11 結果全部是 23",
        "症狀：API 仍可能回 200，但輸出全部 fail-safe 23。",
        "原因：required model failure/timeout、回傳 key 不完整/非有限數值、超過影像門檻，或 publish reserve 到點。",
        "處理：讀 backup/.../ai_result/manifest.json 的 reason/errors/timing，再查 system.inference 與模型 logs。這通常是安全策略正常啟動，不應先改成 optional。",
    ),
    (
        "14.12 SQLite locked / 狀態不可靠",
        "症狀：鎖定錯誤、WAL 異常、worker 看見不同 queue。",
        "原因：database_path 在 SMB/UNC、各 worker 指到不同 config/database，或防毒/備份程式干預。",
        "處理：移到 AIPC 本機 NTFS，確認所有服務的 AI_CONFIG_PATH 一致；排除即時備份掃描。維修前先停止全部 worker。",
    ),
    (
        "14.13 Job 卡在 PUBLISHING 或 PRIMARY_RETURNED",
        "症狀：長時間未 DONE。",
        "原因：PUBLISHING 多為 machine return/SMB 寫入重試；PRIMARY_RETURNED 表示 Primary 已可見，只剩本機 ai_result finalize。",
        "處理：查看 system.publisher、權限、磁碟空間與 share。修復後讓 Publisher 依 lease 自動回復；不要重跑 inference 或手改 Primary。",
    ),
    (
        "14.14 NSSM 服務找不到磁碟、模型或設定",
        "症狀：手動可跑，服務模式卻 FileNotFound/Permission denied。",
        "原因：AppDirectory 未設、相對路徑基準錯、服務帳號無權限、映射磁碟不存在、AppEnvironmentExtra 不完整。",
        "處理：固定 AppDirectory、使用完整 PYTHON_EXE/AI_CONFIG_PATH/model path；SMB 改 UNC 並指定服務帳號。以該帳號做讀寫測試。",
    ),
    (
        "14.15 Primary 寫入失敗或超過 30 秒",
        "症狀：deadline_met=false、share timeout、Publisher 重試。",
        "原因：SMB 不可用、原始備份過慢、GPU/CPU 過載、staging_root 造成第二次複製，或模型超過 publish cutoff。",
        "處理：先修 share 與本機 IO；staging_root 保持 null；檢查 ready_at 到各 stage latency。30 秒是 soft real-time，AIPC/SMB 中斷時無法保證。",
    ),
    (
        "14.16 high cover 規則永遠不觸發",
        "症狀：其他缺陷正常，但沒有 high cover。",
        "原因：paste 8001 預設停用；cover% 還需要 paste_pixels 與由 Width/Length 得到的 pad_area。",
        "處理：只有在確定需求後才離線打包 MobileSAM/paste extra、啟動 8001、enabled=true 並重啟相關程序；驗證新增負載不破壞 30 秒預算。",
    ),
]

for title, symptom, cause, action in troubles:
    add_heading(title, 2)
    add_para(symptom, bold_lead="症狀：", size=10.2, after=3)
    add_para(cause, bold_lead="原因：", size=10.2, after=3)
    add_para(action, bold_lead="處理：", size=10.2, after=7)

add_heading("15. 快速命令表", 1)
add_heading("15.1 版本與環境", 2)
add_code(r"""
%PYTHON_EXE% --version
%PYTHON_EXE% -c "import sys; print(sys.executable); print(sys.version)"
%PYTHON_EXE% scripts\check_env.py
%PYTHON_EXE% -m pip check
%PYTHON_EXE% -c "import torch; print(torch.__version__, torch.version.cuda, torch.cuda.is_available())"
%PYTHON_EXE% -c "import onnxruntime as ort; print(ort.__version__, ort.get_available_providers())"
%PYTHON_EXE% -c "import tensorrt as trt; print(trt.__version__)"
""")
add_heading("15.2 Worker 診斷", 2)
add_code(r"""
01_pipeline_ingest.bat --once --date 2026-08-31
02_pipeline_inference.bat --once --worker-id manual-inference
03_pipeline_publisher.bat --once
""")
add_callout("使用條件", "--once 仍會對設定中的資料庫與目錄採取實際動作。只在隔離測試 config 或確認沒有正式待處理 job 時使用。", kind="warn")
add_heading("15.3 健康檢查", 2)
add_code(r"""
Invoke-RestMethod http://127.0.0.1:8000/health | ConvertTo-Json -Depth 6
Invoke-RestMethod http://127.0.0.1:8002/health | ConvertTo-Json -Depth 6
Invoke-RestMethod http://127.0.0.1:5050/ready  | ConvertTo-Json -Depth 6
""")
add_para("5050 /ready 只屬相容 API，並不代表三個 durable worker 的 heartbeat/queue 都正常；stage 狀態仍需看 SQLite、system.* logs 與 manifest。")

add_heading("16. 官方與專案參考", 1)
sources = [
    ("專案 README.md", "專案根目錄 README.md"),
    ("專案 pyproject.toml / uv.lock", "依賴、Python 範圍、平台 marker 與鎖定版本"),
    ("專案 config/ai_server.example.json", "設定範本"),
    ("Python 3.12 venv", "https://docs.python.org/3.12/library/venv.html"),
    ("uv CLI：offline / frozen / no-python-downloads", "https://docs.astral.sh/uv/reference/cli/"),
    ("uv cache", "https://docs.astral.sh/uv/concepts/cache/"),
    ("pip download", "https://pip.pypa.io/en/stable/cli/pip_download/"),
    ("Conda Pack", "https://conda.github.io/conda-pack/"),
    ("NVIDIA TensorRT Engine Compatibility", "https://docs.nvidia.com/deeplearning/tensorrt/latest/inference-library/engine-compatibility.html"),
    ("NVIDIA TensorRT Support Matrix", "https://docs.nvidia.com/deeplearning/tensorrt/latest/getting-started/support-matrix.html"),
]
for label, target in sources:
    p = doc.add_paragraph()
    set_num(p, bullet_num_id)
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(label + "："), size=9.7, bold=True)
    if urlparse(target).scheme:
        add_hyperlink(p, target, target)
    else:
        set_run_font(p.add_run(target), size=9.7)

add_heading("文件維護建議", 2)
add_bullet("每次更新 Python、CUDA、PyTorch、ONNX Runtime、TensorRT、模型或設定 schema 時，重做離線重建與端到端驗收。")
add_bullet("把 uv.lock 納入正式版本控制或部署產物清單；目前只存在工作區且被 .gitignore 忽略，容易在交付時遺漏。")
add_bullet("將每台 AIPC 的版本矩陣、模型 SHA256、config SHA256、驗收結果與回復包位置納入變更單。")

# Core properties and final save
doc.core_properties.title = "AI SPI Inference API 虛擬環境與離線 AIPC 部署手冊"
doc.core_properties.subject = "Conda、uv、venv、離線部署、設定、服務化與 troubleshooting"
doc.core_properties.author = "OpenAI Codex"
doc.core_properties.keywords = "AI SPI, AIPC, offline deployment, uv, conda, venv, TensorRT"
doc.core_properties.comments = "Generated from the project workspace on 2026-08-31."

DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(DOCX_OUT)
print(DOCX_OUT)
